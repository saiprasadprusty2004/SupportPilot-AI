import hashlib
import io
import math
import re
import zipfile
from functools import lru_cache
from threading import Lock

from sqlalchemy import delete, select, text

from app.config import settings
from app.models import Chunk, Document
from app.services.storage import storage

MIME = {
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".txt": "text/plain",
    ".md": "text/markdown",
}
_MODEL = None
_MODEL_LOCK = Lock()


@lru_cache(maxsize=1)
def model():
    global _MODEL
    from sentence_transformers import SentenceTransformer

    with _MODEL_LOCK:
        if _MODEL is None:
            _MODEL = SentenceTransformer(
                settings().embedding_model,
                revision=settings().embedding_revision,
                device="cpu",
                trust_remote_code=False,
            )
        return _MODEL


def embed(texts):
    if settings().demo_mode:
        # Deterministic lexical fixture only. Never used in live mode or described as semantic embeddings.
        vectors = []
        for value in texts:
            vector = [0.0] * 384
            for word in re.findall(r"\w+", value.lower()):
                position = int(hashlib.sha256(word.encode()).hexdigest()[:8], 16) % 384
                vector[position] += 1
            norm = math.sqrt(sum(x * x for x in vector)) or 1
            vectors.append([x / norm for x in vector])
        return vectors
    return model().encode(texts, normalize_embeddings=True, batch_size=32).tolist()


def split_text(value):
    if settings().demo_mode:
        tokens = value.split()
        return [" ".join(tokens[i : i + 200]) for i in range(0, len(tokens), 168)]
    tokenizer = model().tokenizer
    tokens = tokenizer.encode(value, add_special_tokens=False)
    return [
        tokenizer.decode(tokens[i : i + 200], skip_special_tokens=True) for i in range(0, len(tokens), 168)
    ]


def extract(data, mime):
    if mime == "application/pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted or len(reader.pages) > 500:
            raise ValueError("PDF is encrypted or exceeds 500 pages")
        pages = [(i + 1, page.extract_text() or "") for i, page in enumerate(reader.pages)]
    elif mime == MIME[".docx"]:
        from docx import Document as WordDocument

        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 100 * 1024 * 1024:
                raise ValueError("DOCX expanded size exceeds 100 MB")
        doc = WordDocument(io.BytesIO(data))
        lines = [p.text for p in doc.paragraphs]
        lines += [" | ".join(cell.text for cell in row.cells) for table in doc.tables for row in table.rows]
        pages = [(1, "\n".join(lines))]
    else:
        pages = [(1, data.decode("utf-8-sig"))]
    total = sum(len(value) for _, value in pages)
    if total > 2_000_000:
        raise ValueError("Extracted text exceeds 2 million characters")
    if total < 10:
        raise ValueError("No searchable text found. Scanned PDFs require OCR, which is not included.")
    return pages


def index_document(db, doc: Document):
    if doc.deleted_at:
        return
    doc.status, doc.error = "indexing", None
    db.commit()
    try:
        data = storage().get(doc.storage_key)
        pieces = [
            (page, chunk)
            for page, value in extract(data, doc.mime_type)
            for chunk in split_text(value)
            if chunk.strip()
        ]
        if len(pieces) > 15000:
            raise ValueError("Document exceeds 15,000 chunks")
        vectors = embed([chunk for _, chunk in pieces])
        db.refresh(doc)
        if doc.deleted_at:
            return
        db.execute(delete(Chunk).where(Chunk.document_id == doc.id, Chunk.workspace_id == doc.workspace_id))
        for (page, value), vector in zip(pieces, vectors):
            db.add(
                Chunk(
                    workspace_id=doc.workspace_id,
                    document_id=doc.id,
                    version=doc.version,
                    text=value,
                    page=page,
                    embedding=vector,
                )
            )
        doc.status = "ready"
        doc.embedding_config = {
            "model": "demo-lexical-fixture" if settings().demo_mode else settings().embedding_model,
            "revision": settings().embedding_revision,
            "dimensions": 384,
            "chunk_tokens": 200,
            "overlap": 32,
        }
        db.commit()
    except Exception as exc:
        db.rollback()
        doc.status, doc.error = "failed", str(exc)[:1000]
        db.commit()
        raise


def search(db, workspace_id, query, limit=6):
    vector = embed([query])[0]
    if db.bind.dialect.name == "postgresql":
        # Both branches filter authorization and current document version before rank fusion.
        db.execute(text("SET LOCAL hnsw.iterative_scan = 'strict_order'"))
        rows = db.execute(
            text("""
          WITH vector_results AS MATERIALIZED (
            SELECT c.id, row_number() OVER (ORDER BY embedding <=> CAST(:vector AS vector)) AS rank
            FROM chunks c JOIN documents d ON d.id=c.document_id AND d.workspace_id=c.workspace_id
            WHERE c.workspace_id=:workspace AND d.deleted_at IS NULL AND d.status='ready'
              AND c.version=d.version
            ORDER BY embedding <=> CAST(:vector AS vector) LIMIT 40
          ), keyword_results AS MATERIALIZED (
            SELECT c.id, row_number() OVER (ORDER BY ts_rank_cd(to_tsvector('english', c.text),
              plainto_tsquery('english', :query)) DESC) AS rank
            FROM chunks c JOIN documents d ON d.id=c.document_id AND d.workspace_id=c.workspace_id
            WHERE c.workspace_id=:workspace AND d.deleted_at IS NULL AND d.status='ready'
              AND c.version=d.version AND to_tsvector('english', c.text) @@ plainto_tsquery('english', :query)
            ORDER BY ts_rank_cd(to_tsvector('english', c.text), plainto_tsquery('english', :query)) DESC LIMIT 40
          ), ranks AS (
            SELECT coalesce(v.id,k.id) AS id, coalesce(1.0/(60+v.rank),0) + coalesce(1.0/(60+k.rank),0) AS score
            FROM vector_results v FULL OUTER JOIN keyword_results k ON v.id=k.id
          ) SELECT c.*,d.name,d.source_url,r.score FROM ranks r JOIN chunks c ON c.id=r.id
              JOIN documents d ON d.id=c.document_id
            WHERE c.workspace_id=:workspace AND 1 - (c.embedding <=> CAST(:vector AS vector)) > 0.15
            ORDER BY r.score DESC LIMIT :limit
        """),
            {"vector": str(vector), "query": query, "workspace": workspace_id, "limit": limit},
        ).mappings()
        return [
            {
                "chunk_id": r["id"],
                "document_id": r["document_id"],
                "version": r["version"],
                "name": r["name"],
                "page": r["page"],
                "excerpt": r["text"],
                "score": float(r["score"]),
                "url": f"/api/v1/documents/{r['document_id']}/source?version={r['version']}",
            }
            for r in rows
        ]
    rows = db.execute(
        select(Chunk, Document)
        .join(Document, Chunk.document_id == Document.id)
        .where(
            Chunk.workspace_id == workspace_id,
            Document.workspace_id == workspace_id,
            Document.deleted_at.is_(None),
            Document.status == "ready",
            Chunk.version == Document.version,
        )
    ).all()
    scored = [(sum(a * b for a, b in zip(vector, chunk.embedding)), chunk, doc) for chunk, doc in rows]
    scored.sort(key=lambda value: value[0], reverse=True)
    return [
        {
            "chunk_id": c.id,
            "document_id": d.id,
            "version": c.version,
            "name": d.name,
            "page": c.page,
            "excerpt": c.text,
            "score": score,
            "url": f"/api/v1/documents/{d.id}/source?version={c.version}",
        }
        for score, c, d in scored[:limit]
        if score > 0.05
    ]
